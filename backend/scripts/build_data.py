"""Run once (or whenever a college .docx changes):  python scripts/build_data.py
Reads ./source_docs/*.docx and writes ./data/*.json + copies the files to ./docs/official/ for download."""
import re, json, shutil
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent          # backend/
SRC, OUT, DOCS = ROOT / 'source_docs', ROOT / 'data', ROOT / 'docs' / 'official'
OUT.mkdir(exist_ok=True); DOCS.mkdir(parents=True, exist_ok=True)
MO = {m: i for i, m in enumerate('January February March April May June July August September October November December'.split(), 1)}
clean = lambda s: re.sub(r'\s+', ' ', s).strip()
fix = lambda s: s.replace('Shasikant', 'Shrikant')          # spelling correction requested by user


def body(doc):
    for el in doc.element.body.iterchildren():
        if el.tag == qn('w:p'): yield 'p', Paragraph(el, doc)
        elif el.tag == qn('w:tbl'): yield 't', Table(el, doc)


def rows(t): return [[clean(c.text) for c in r.cells] for r in t.rows]   # r.cells is grid-aligned (merged cells repeat)


def resolve(raw, courses):
    raw = re.sub(r'\s*/\s*', '/', raw).strip()
    if len(raw) <= 1 or not raw: return None                  # empty or LUNCH column letter
    if raw.upper() == 'LIB': return {'k': 'LIB'}
    p = raw.split('/')
    if len(p) < 3: return {'k': raw, 's': raw, 'f': '', 'r': '', 't': ''}
    t, a, f = p[:3]; room = '/'.join(p[3:])
    meta = lambda c: c[3].replace(' ', '').split('/')
    hit = next((c for c in courses if meta(c)[:3] == [t, a, f]), None) \
       or next((c for c in courses if len(meta(c)) > 2 and meta(c)[0] == t and meta(c)[2] == f), None) \
       or next((c for c in courses if len(meta(c)) > 2 and meta(c)[1] == a and meta(c)[2] == f), None)
    return {'k': a, 't': t, 's': hit[2] if hit else a, 'f': hit[4] if hit else '', 'r': room}


def timetables(path):
    out, pend = {}, None
    cur = None
    for k, o in body(Document(path)):
        if k == 'p':
            m = re.findall(r'Section:\s*([A-Za-z()\-0-9]+)', o.text)
            if not m: continue
            if cur is not None and cur['sec'] is None: cur['sec'] = m[-1]
            else: pend = m[-1]
        else:
            r = rows(o)
            if r[0][0].startswith(('Day', 'Time')):
                cur = {'sec': pend, 'raw': {x[0]: (x[1:9] + [''] * 8)[:8] for x in r[1:]}, 'courses': []}; pend = None
                out[id(cur)] = cur
            elif r[0][0].startswith('Credit') and cur: cur['courses'] = r[1:]
    res = {}
    for c in out.values():
        sec = c['sec']
        res[sec] = {'slots': {d: [resolve(x, c['courses']) for x in v] for d, v in c['raw'].items()},
                    'courses': [{'code': x[1], 'name': x[2], 'f': x[4]} for x in c['courses'] if x[2]]}
    return res


def pdate(s):
    s = clean(s).replace('Dushehra', 'Dussehra')
    if not re.search(r'[A-Z][a-z]+ \d', s) or s.startswith('First'): return None, None
    y = re.search(r'\d{4}', s).group(); parts = re.split(r'\s+to\s+', s)
    def one(p, mo=None):
        m = re.match(r'(?:([A-Za-z]+) )?(\d+)', p.strip()); return (m.group(1) or mo, int(m.group(2)))
    a = one(parts[0]); b = one(parts[1], a[0]) if len(parts) > 1 else None
    iso = lambda x: f'{y}-{MO[x[0]]:02d}-{x[1]:02d}'
    return iso(a), iso(b) if b else None


def table_after(path, first_cell):
    for k, o in body(Document(path)):
        if k == 't':
            r = rows(o)
            if r[0][0].lower().startswith(first_cell): return r[1:]
    return []


def syllabus(path):
    res, name, on = {}, None, False
    for k, o in body(Document(path)):
        if k == 'p':
            st, tx = o.style.name, clean(o.text)
            if st.startswith('Heading 1'): name, on = tx, False
            elif st.startswith('Heading 2'): on = tx.lower().startswith('course contents')
            elif on and tx and name: res[name] = res.get(name, '') + tx + '\n'
        elif on and name:
            for r in rows(o): res[name] = res.get(name, '') + ' | '.join(dict.fromkeys(x for x in r if x)) + '\n'
    return {k: v.strip() for k, v in res.items()}


f = lambda n: next(SRC.glob(n))
tt = {}
for p in (f('CS-3H*.docx'), f('Time_Table_2nd*.docx')): tt.update(timetables(p))
cal = []
for r in table_after(f('Academic_Calendar*.docx'), 's.'):
    s, e = pdate(r[2]); cal.append({'n': r[1], 's': s, 'e': e, 'raw': r[2]})
hol = []
for r in table_after(f('Holidays*.docx'), 's.'):
    s, e = pdate(r[3]); hol.append({'n': r[1].lstrip('*'), 's': s, 'e': e})
coord = {re.sub(r'[^A-Z0-9]', '', r[1].upper()): {'sec': r[1], 'name': r[2], 'mob': r[3]} for r in table_after(f('List_of_Class*.docx'), 'sl.')}
for name, obj in (('tt', tt), ('cal', cal), ('hol', hol), ('coord', coord), ('syl', syllabus(f('5th_Semester*.docx')))):
    (OUT / f'{name}.json').write_text(fix(json.dumps(obj, ensure_ascii=False, indent=1)), encoding='utf-8')
for p in SRC.glob('*.docx'): shutil.copy(p, DOCS / p.name)
print('sections:', list(tt)); print('coordinators:', len(coord), 'calendar rows:', len(cal), 'holidays:', len(hol))
